import streamlit as st
import sqlite3
import pandas as pd
import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Alignment, Font
import matplotlib.pyplot as plt

DB_PATH = "diario_bordo.db"

# =====================================================
# BANCO DE DADOS - GARANTIR QUE A TABELA EXISTE
# =====================================================
def init_db():
    conn = sqlite3.connect(DB_PATH)
    return conn

def create_table_if_not_exists(conn):
    try:
        conn.execute("""
        CREATE TABLE IF NOT EXISTS diario_bordo (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            motorista TEXT,
            placa TEXT,
            data TEXT,
            viagem TEXT,
            entrega TEXT,
            cliente TEXT,
            cidade TEXT,
            hora_primato TEXT,
            hora_cliente TEXT,
            nota_fiscal TEXT,
            peso REAL,
            tipo_carga TEXT,
            km REAL,
            km_aves REAL,
            km_suino REAL,
            km_bovino REAL,
            km_farinha_carne REAL,
            checklist TEXT,
            observacoes TEXT
        );
        """)
        conn.commit()
    except Exception as e:
        st.error(f"Erro ao criar tabela: {e}")

# =====================================================
# Funções padrão
# =====================================================
def insert_entry(conn, data):
    cols = ",".join(data.keys())
    placeholders = ",".join("?" for _ in data)
    sql = f"INSERT INTO diario_bordo ({cols}) VALUES ({placeholders})"
    conn.execute(sql, list(data.values()))
    conn.commit()

def query_entries(conn, start_date=None, end_date=None, motorista=None, placa=None, tipo_carga=None, cliente=None, cidade=None):
    sql = "SELECT * FROM diario_bordo WHERE 1=1"
    params = []

    if start_date:
        sql += " AND date(data) >= date(?)"
        params.append(start_date)

    if end_date:
        sql += " AND date(data) <= date(?)"
        params.append(end_date)

    if motorista:
        if isinstance(motorista, (list,tuple)):
            sql += " AND motorista IN ({})".format(",".join("?" for _ in motorista))
            params.extend(motorista)
        else:
            sql += " AND motorista = ?"
            params.append(motorista)

    if placa:
        if isinstance(placa, (list,tuple)):
            sql += " AND placa IN ({})".format(",".join("?" for _ in placa))
            params.extend(placa)
        else:
            sql += " AND placa = ?"
            params.append(placa)

    if tipo_carga:
        if isinstance(tipo_carga, (list,tuple)):
            sql += " AND tipo_carga IN ({})".format(",".join("?" for _ in tipo_carga))
            params.extend(tipo_carga)
        else:
            sql += " AND tipo_carga = ?"
            params.append(tipo_carga)

    if cliente:
        sql += " AND cliente LIKE ?"
        params.append(f"%{cliente}%")

    if cidade:
        sql += " AND cidade LIKE ?"
        params.append(f"%{cidade}%")

    df = pd.read_sql_query(sql, conn, params=params)
    return df

def delete_entry(conn, entry_id):
    conn.execute("DELETE FROM diario_bordo WHERE id = ?", (entry_id,))
    conn.commit()

# =====================================================
# Excel formatado
# =====================================================
def gerar_excel_formatado(df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"

    header_font = Font(bold=True)

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value)) for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max_len

    from io import BytesIO
    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    return arquivo

# =====================================================
# DADOS FIXOS
# =====================================================
motoristas_lista = [
    "Andre", "Paulo", "Paulo Mandotti", "Jacinto",
    "Jeferson", "Edson", "Rosenildo", "Alan", "Fabio"
]

placas_opcoes = [
    "BCK0I38", "BCK5D22", "BCI5J96", "BCI9J76",
    "BCI2B70", "BCI9J78", "BCJ6I67", "BCJ6856", "RHP6C77"
]

# =====================================================
# INTERFACE
# =====================================================
st.set_page_config(page_title="Diário de Bordo - Primato", layout="wide")

# BANCO DEVE SER CRIADO AQUI
conn = init_db()
create_table_if_not_exists(conn)
conn.close()

# LOGO
st.image("logo_primato.png", width=200)
st.title("Sistema de Diário de Bordo - Primato")

# LOGIN
modo = st.radio("Selecione o modo de acesso:", ["Motorista", "Administrador"])

if "logado" not in st.session_state:
    st.session_state.logado = False
if "motorista" not in st.session_state:
    st.session_state.motorista = None
if "admin" not in st.session_state:
    st.session_state.admin = False

# --------------- LOGIN MOTORISTA --------------------
if not st.session_state.logado and modo == "Motorista":
    with st.form("login_motorista"):
        nome = st.selectbox("Motorista:", motoristas_lista)
        senha = st.text_input("Senha (primeiro nome)", type="password")
        ok = st.form_submit_button("Entrar")

        if ok:
            if senha.lower() == nome.split()[0].lower():
                st.session_state.logado = True
                st.session_state.motorista = nome
                st.rerun()
            else:
                st.error("Senha incorreta!")

# --------------- LOGIN ADMIN --------------------
if modo == "Administrador" and not st.session_state.admin:
    with st.form("login_admin"):
        user = st.text_input("Usuário")
        senha = st.text_input("Senha", type="password")
        ok = st.form_submit_button("Entrar")

        if ok:
            if user == "ADMIN" and senha == "primato2025":
                st.session_state.admin = True
                st.rerun()
            else:
                st.error("Credenciais inválidas")

# =====================================================
# PAINEL DO MOTORISTA
# =====================================================
if st.session_state.logado:

    st.header(f"Motorista: {st.session_state.motorista}")

    conn = init_db()
    create_table_if_not_exists(conn)

    aba1, aba2 = st.tabs(["Registrar Viagem", "Relatórios"])

    with aba1:
        st.subheader("Registrar nova viagem")

        with st.form("registro", clear_on_submit=True):
            placa = st.selectbox("Placa", placas_opcoes)
            data = st.date_input("Data", datetime.today())
            viagem = st.text_input("Viagem (nº)")
            entrega = st.text_input("Entrega")
            cliente = st.text_input("Cliente")
            cidade = st.text_input("Cidade")
            hora_primato = st.text_input("Hora Primato")
            hora_cliente = st.text_input("Hora Cliente")
            nota = st.text_input("Nota Fiscal")
            peso = st.number_input("Peso (kg)", 0.0)
            tipo = st.selectbox("Tipo de carga", ["AVES","SUÍNO","BOVINO","FARINHA_DE_CARNE","PEIXE","OUTRO"])
            km = st.number_input("KM total", 0.0)
            obs = st.text_area("Observações")

            enviar = st.form_submit_button("Salvar")

            if enviar:
                insert_entry(conn, {
                    "motorista": st.session_state.motorista,
                    "placa": placa,
                    "data": data.strftime("%Y-%m-%d"),
                    "viagem": viagem,
                    "entrega": entrega,
                    "cliente": cliente,
                    "cidade": cidade,
                    "hora_primato": hora_primato,
                    "hora_cliente": hora_cliente,
                    "nota_fiscal": nota,
                    "peso": peso,
                    "tipo_carga": tipo,
                    "km": km,
                    "km_aves": 0,
                    "km_suino": 0,
                    "km_bovino": 0,
                    "km_farinha_carne": 0,
                    "checklist": json.dumps({}),
                    "observacoes": obs
                })
                st.success("Registro salvo!")

    with aba2:
        st.subheader("Relatórios do motorista")

        df = query_entries(conn, motorista=st.session_state.motorista)

        if df.empty:
            st.info("Nenhum registro encontrado.")
        else:
            st.dataframe(df)

            excel = gerar_excel_formatado(df)
            st.download_button("📄 Baixar Excel", data=excel,
                               file_name="relatorio.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    st.button("Sair", on_click=lambda: st.session_state.clear())

    conn.close()

# =====================================================
# PAINEL DO ADMINISTRADOR
# =====================================================
if st.session_state.admin:

    st.header("Painel Administrativo")
    conn = init_db()
    create_table_if_not_exists(conn)

    df = query_entries(conn)

    if df.empty:
        st.info("Nenhum registro encontrado.")
    else:
        st.dataframe(df)

        # Excel
        excel = gerar_excel_formatado(df)
        st.download_button("📄 Exportar Excel", data=excel,
                           file_name="relatorio_admin.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        # Deletar
        ids = st.multiselect("IDs para excluir:", df["id"])
        if st.button("Excluir selecionados"):
            for i in ids:
                delete_entry(conn, i)
            st.success("Registros excluídos.")
            st.rerun()

    st.button("Logout", on_click=lambda: st.session_state.clear())

    conn.close()
